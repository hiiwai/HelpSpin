"""Writing a matplotlib figure to disk, shared by every canvas.

The main spectrum canvas and the CEST plots offer the same "Save image"
formats and must produce the same kind of file: one implementation, so a fix
in one place (a format, editable text in vector output) reaches both.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

# (file-dialog filter, extension). EMF is deliberately absent: matplotlib
# cannot write it, and the only route (an external Inkscape) silently
# produced nothing when Inkscape was missing.
EXPORT_FILTERS = [
    ("PNG image (*.png)", "png"),
    ("JPEG image (*.jpg)", "jpg"),
    ("TIFF image (*.tif)", "tif"),
    ("SVG vector (*.svg)", "svg"),
    ("PDF vector (*.pdf)", "pdf"),
    ("EPS vector (*.eps)", "eps"),
    ("PostScript (*.ps)", "ps"),
    ("PowerPoint slide (*.pptx)", "pptx"),
]

EXPORT_SUFFIXES = {ext for _label, ext in EXPORT_FILTERS} | {"jpeg", "tiff"}


def suffix_for_filter(chosen: str, default: str = "png") -> str:
    """Extension belonging to a file-dialog filter string."""
    return next((ext for label, ext in EXPORT_FILTERS if label == chosen), default)


def save_figure(figure, path, *, transparent: bool = False, dpi: int = 300) -> None:
    """Write `figure` exactly as displayed; the format follows the suffix.

    No tight bounding box and no layout change: those re-crop and re-scale,
    so the file would not match the canvas. Vector formats keep text as
    editable text rather than outlines.
    """
    import matplotlib as mpl

    path = str(path)
    # Path().suffix, never a split on ".": a dot in a DIRECTORY name would
    # otherwise be read as the format.
    suffix = Path(path).suffix.lstrip(".").lower() or "png"
    if suffix == "pptx":
        _save_pptx(figure, path, dpi=dpi)
        return
    overrides = {"svg.fonttype": "none", "pdf.fonttype": 42, "ps.fonttype": 42}
    with mpl.rc_context(overrides):
        figure.savefig(
            path, format=suffix, dpi=dpi,
            facecolor="none" if transparent else "white",
            edgecolor="none", transparent=transparent,
        )


def _save_pptx(figure, path: str, dpi: int = 300) -> None:
    """One slide sized to the figure, holding a full-resolution PNG of it."""
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:   # pragma: no cover - dependency present
        raise RuntimeError("PowerPoint export needs the python-pptx package") from exc

    width_in, height_in = figure.get_size_inches()
    with tempfile.TemporaryDirectory() as tmp:
        png = f"{tmp}/slide.png"
        figure.savefig(png, format="png", dpi=dpi, facecolor="white",
                       edgecolor="none")
        prs = Presentation()
        prs.slide_width = Inches(float(width_in))
        prs.slide_height = Inches(float(height_in))
        slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank layout
        slide.shapes.add_picture(png, 0, 0, width=Inches(float(width_in)),
                                 height=Inches(float(height_in)))
        prs.save(path)
